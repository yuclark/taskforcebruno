import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc_path = "docs/Team2526-sem2-it332-18_WildcatWelfare_SDD-3 (1).docx"
doc = docx.Document(doc_path)

print(f"Initial paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}")

# Helper to style a cell
def style_cell(cell, bg_color=None, bold=False, italic=False, font_size=9.5):
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

# Helper to add styled paragraph after another
def add_paragraph_after(target_p, text="", style_name="normal"):
    new_p = doc.add_paragraph(text, style=style_name)
    target_p._p.addnext(new_p._p)
    return new_p

# 1. UPDATE CHANGE HISTORY
# Look for paragraph with 'Change History Signature'
ch_idx = None
for i, p in enumerate(doc.paragraphs):
    if "Change History Signature" in p.text:
        ch_idx = i
        break

if ch_idx is not None:
    print(f"Found Change History at P{ch_idx}")
    target_p = doc.paragraphs[ch_idx]
    # Create table after target_p
    # In docx, to insert a table at a specific location:
    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    target_p._p.addnext(tbl._tbl)

    headers = ["Version", "Date", "Author", "Summary of Changes / Refactoring Description"]
    col_widths = [Inches(0.9), Inches(1.2), Inches(1.5), Inches(3.4)]

    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.width = col_widths[j]
        style_cell(cell, bg_color="EAEAEA", bold=True)

    history_data = [
        ("1.0", "February 2026", "Wildcat Welfare Team", "Initial Software Design Description Draft & Scope Definition."),
        ("2.0", "May 2026", "Development Team", "Detailed Component Decomposition, Class Diagrams, and Sequence Flows."),
        ("3.0", "August 2026", "Development Team", "Pre-Validation Baseline Architecture & Database Relational Schemas."),
        ("3.1", "September 19, 2026", "Task Force Bruno Team", "Post-MVP Validation Refactoring: Incorporated empirical findings from N=34 campus evaluation. Added Public Emergency Pet Passport (unauthenticated QR access), 14-day bite/scratch quarantine telemetry (RA 9482), 1-click sighting intake handoff, printable collar QR generator, dimensional supply logistics (kg/cans/vials), and interactive user onboarding.")
    ]

    for row_idx, data in enumerate(history_data, start=1):
        bg = "FAFAFA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = tbl.cell(row_idx, col_idx)
            cell.text = val
            cell.width = col_widths[col_idx]
            style_cell(cell, bg_color=bg, bold=(col_idx == 0))

print("Change History Table created successfully.")

# 2. UPDATE PREFACE PILLARS
for p in doc.paragraphs:
    if "Instant Identification:" in p.text:
        p.text = "Instant Identification & Public Pet Passport: QR-coded profiles for real-time access to companion health history, featuring an unauthenticated emergency public passport view for immediate safety and rabies verification without mandatory login."
    elif "Clinical Governance:" in p.text:
        p.text = "Clinical Governance & Biosecurity Telemetry: A secure MDC Dashboard for logging medical incidents, vaccination schedules, and active 14-day observation quarantine tracking for bite/scratch incidents in compliance with RA 9482."
    elif "Campus Vigilance:" in p.text:
        p.text = "Campus Vigilance & 1-Click Sighting Intake: A community reporting hub with CIT-U landmark localization and an administrative 1-click handoff directly admitting stray sightings into pet registration."
    elif "Humane Rehoming:" in p.text:
        p.text = "Humane Rehoming: A digital adoption pipeline tracking the end-to-end journey of rehoming pets with living environment validation and clear recovery/adoptability badges."
    elif "Resource Transparency:" in p.text:
        p.text = "Resource & Supply Logistics: A live inventory tracker monitoring medical and nutritional supplies by genuine physical units (kilograms, cans, vials, pieces) and generating printable collar QR tags."

print("Preface pillars updated.")

# 3. UPDATE SCOPE
for p in doc.paragraphs:
    if "QR Coded Pet Profiles:" in p.text:
        p.text = "QR Coded Pet Profiles & Public Emergency Passport: Digitized identification system providing instant access to companion vaccination records, health history, and emergency clinic/security contacts via camera QR scanning without requiring an institutional login."
    elif "MDC Medical Dashboard:" in p.text:
        p.text = "MDC Medical Dashboard & Biosecurity Surveillance: Restricted to clinic staff to log, manage, and track clinical visits, automated vaccination reminders, and active 14-day observation quarantine alerts for animal bite or scratch incidents under Republic Act No. 9482."
    elif "Community Newsfeed and Reporting Hub:" in p.text:
        p.text = "Community Newsfeed & Sighting Triage Hub: A live module for crowdsourced reporting of stray sightings with CIT-U landmark localization, missing pet alerts, and administrative 1-click admission into pet registration."
    elif "Digital Adoption Pipeline:" in p.text:
        p.text = "Digital Adoption Pipeline: An integrated adoption application system that enables end-to-end vetting, living condition verification, and transition to Adopted Alumni status."
    elif "Resource and Donation Tracker:" in p.text:
        p.text = "Resource & Supply Logistics Tracker: Live inventory dashboard monitoring medical and nutritional supplies by genuine physical units (kg, cans, vials, assets) with printable collar QR tag generation."

print("Scope updated.")

# 4. MODULE 1.1: ADD Public Pet Passport & Camera Permission Fallback
m1_target = None
for i, p in enumerate(doc.paragraphs):
    if "PetProfileView" in p.text and "React/UI Component" in doc.paragraphs[i+2].text:
        m1_target = i + 2
        break

if m1_target:
    print(f"Found PetProfileView at P{m1_target}")
    tp = doc.paragraphs[m1_target]
    
    # Add PetPassportView
    p1 = add_paragraph_after(tp, "PetPassportView", style_name="normal")
    p1.runs[0].font.bold = True
    p2 = add_paragraph_after(p1, "A public-facing, unauthenticated emergency passport interface rendered when campus community members or visitors scan an animal's collar QR tag. It displays verified rabies immunization badges, emergency medical alerts, and direct campus clinic/security hotlines without requiring an initial institutional login.", style_name="normal")
    p3 = add_paragraph_after(p2, "React / Public Emergency UI Component", style_name="normal")
    p3.runs[0].font.italic = True
    
    # Add QRScannerView enhancement
    p4 = add_paragraph_after(p3, "QRScannerView (Hardware Camera & File Upload Fallback)", style_name="normal")
    p4.runs[0].font.bold = True
    p5 = add_paragraph_after(p4, "The optical scanning interface utilizing jsQR. Includes environmental rear-camera constraints and an automated gallery file-upload fallback when hardware camera permissions are restricted or denied on mobile devices.", style_name="normal")
    p6 = add_paragraph_after(p5, "React / Optical Decoder Component", style_name="normal")
    p6.runs[0].font.italic = True

# Also add PublicPassportAPIController to Module 1.1 Backend
m1_be_target = None
for i, p in enumerate(doc.paragraphs):
    if "PetDataService" in p.text and "Service Layer" in doc.paragraphs[i+2].text:
        m1_be_target = i + 2
        break

if m1_be_target:
    print(f"Found PetDataService at P{m1_be_target}")
    tp = doc.paragraphs[m1_be_target]
    p1 = add_paragraph_after(tp, "PublicPassportAPIController", style_name="normal")
    p1.runs[0].font.bold = True
    p2 = add_paragraph_after(p1, "Exposes a public read-only endpoint (GET /api/pets/passport/:id) that delivers sanitized public safety metadata (name, vaccination tier, rabies status, clinic hotline) while stripping internal administrative logs.", style_name="normal")
    p3 = add_paragraph_after(p2, "Django REST Framework Public API Endpoint", style_name="normal")
    p3.runs[0].font.italic = True

# 5. MODULE 1.2: ADD 14-Day Bite/Scratch Quarantine Tracker (RA 9482)
m12_target = None
for i, p in enumerate(doc.paragraphs):
    if "IncidentReportForm" in p.text and "React Form Component" in doc.paragraphs[i+2].text:
        m12_target = i + 2
        break

if m12_target:
    print(f"Found IncidentReportForm at P{m12_target}")
    tp = doc.paragraphs[m12_target]
    p1 = add_paragraph_after(tp, "QuarantineTrackerComponent", style_name="normal")
    p1.runs[0].font.bold = True
    p2 = add_paragraph_after(p1, "A biosecurity monitoring widget integrated into the MDC Dashboard. Automatically calculates and displays active 14-day observation countdowns for campus animals involved in reported human bite or scratch incidents, enforcing strict regulatory adherence to Republic Act No. 9482 (Anti-Rabies Act of 2007).", style_name="normal")
    p3 = add_paragraph_after(p2, "React Biosecurity Telemetry Component", style_name="normal")
    p3.runs[0].font.italic = True

# Update IncidentReportForm description to mention RA 9482
for p in doc.paragraphs:
    if "A specialized form for reporting bite or scratch incidents" in p.text:
        p.text = "A specialized form for reporting bite or scratch incidents involving campus animals. Captures incident type (Bite/Scratch—Human), incident date, campus landmark location, victim details, and automatically initializes an active 14-day observation quarantine status in compliance with Republic Act No. 9482. Submitted records trigger emergency clinical follow-ups and Post-Exposure Prophylaxis (PEP) tracking."

# 6. MODULE 2.2: ADD Printable Collar Tag QR Generator & Physical Units
m22_target = None
for i, p in enumerate(doc.paragraphs):
    if "Transaction Ledger History" in p.text and "Inline React UI Block" in doc.paragraphs[i+2].text:
        m22_target = i + 2
        break

if m22_target:
    print(f"Found Transaction Ledger History at P{m22_target}")
    tp = doc.paragraphs[m22_target]
    p1 = add_paragraph_after(tp, "CollarTagGenerator", style_name="normal")
    p1.runs[0].font.bold = True
    p2 = add_paragraph_after(p1, "An integrated administrative utility that compiles pet identification codes and URLs into high-resolution, printable collar QR tags formatted for standard waterproof collar tag matrices.", style_name="normal")
    p3 = add_paragraph_after(p2, "React / HTML5 Canvas Tag Export Utility", style_name="normal")
    p3.runs[0].font.italic = True

# Update Active Supply Stock Master & InventorySchema description for physical units
for p in doc.paragraphs:
    if "A horizontal-scrolling tabular interface that lists current items with real-time descriptions" in p.text:
        p.text = "A dynamic tabular interface that lists current items segmented by standardized physical operational units (kilograms for kibble, cans for wet food, vials for vaccines/medications, and pieces for collar/hardware assets). Enforces discrete numerical step increments to prevent inventory data entry errors."
    elif "The structural data definition for the main supply entries." in p.text:
        p.text = "The structural relational data definition for supply entries. Models operational fields including item_name, category, physical unit_type (kilograms, cans, vials, pieces), quantity, min_threshold, donor_attribution, and expiration_date."

# 7. MODULE 3.1: ADD Sighting Triage & User Guide Modal
m31_target = None
for i, p in enumerate(doc.paragraphs):
    if "NewsfeedView" in p.text and "React Functional Component" in doc.paragraphs[i+2].text:
        m31_target = i + 2
        break

if m31_target:
    print(f"Found NewsfeedView at P{m31_target}")
    tp = doc.paragraphs[m31_target]
    
    # SightingTriageView
    p1 = add_paragraph_after(tp, "SightingTriageView", style_name="normal")
    p1.runs[0].font.bold = True
    p2 = add_paragraph_after(p1, "A specialized administrative triage queue allowing campus clinic staff and welfare officers to review incoming community sighting reports, inspect uploaded photos and location landmarks, and execute a 1-click 'Admit & Intake Stray' workflow that auto-populates the new pet registration form without duplicate data entry.", style_name="normal")
    p3 = add_paragraph_after(p2, "React Administrative Triage Grid Component", style_name="normal")
    p3.runs[0].font.italic = True
    
    # UserGuideModal
    p4 = add_paragraph_after(p3, "UserGuideModal (UserGuideModal.jsx)", style_name="normal")
    p4.runs[0].font.bold = True
    p5 = add_paragraph_after(p4, "An interactive onboarding modal designed in response to MVP validation user feedback. Provides first-time students, faculty, and administrative users with step-by-step guidance on QR scanning, sighting submissions, and pet passport navigation.", style_name="normal")
    p6 = add_paragraph_after(p5, "React Interactive Modal / Walkthrough Component", style_name="normal")
    p6.runs[0].font.italic = True

    # ReportSightingView with landmarks
    p7 = add_paragraph_after(p6, "ReportSightingView (with CIT-U Campus Landmark Selector)", style_name="normal")
    p7.runs[0].font.bold = True
    p8 = add_paragraph_after(p7, "The community sighting submission interface. Features structured dropdown selection for standardized CIT-U campus buildings and zones (NGE Building, Main Canteen, Academic Complex, Gymnasium, Grounds) alongside photo attachment and condition tags, eliminating ambiguous location descriptions.", style_name="normal")
    p9 = add_paragraph_after(p8, "React Form / Sighting Submission Component", style_name="normal")
    p9.runs[0].font.italic = True

print(f"Saving updated SDD to {doc_path}...")
doc.save(doc_path)
print("Successfully saved updated SDD docx!")
